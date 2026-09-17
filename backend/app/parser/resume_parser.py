import re
import fitz  # PyMuPDF
import docx
from typing import Dict, List, Optional

try:
    import spacy

    _nlp = spacy.load("en_core_web_sm")
except Exception:
    _nlp = None

def extract_text_from_pdf(path: str) -> str:
    text = ""
    with fitz.open(path) as doc:
        for page in doc:
            text += page.get_text()
    return text

def extract_text_from_docx(path: str) -> str:
    doc = docx.Document(path)
    parts = [para.text for para in doc.paragraphs if para.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)

def extract_text(path: str) -> str:
    lower = path.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(path)
    elif lower.endswith((".docx", ".doc")):
        return extract_text_from_docx(path)
    raise ValueError(f"Unsupported file type: {path}")

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s-]?)?(?:\d[\s-]?){9,11}\d")
NAME_WORD_RE = re.compile(r"^[A-Za-z][A-Za-z'\-\.]*$")
NAME_TITLE_WORD_RE = re.compile(r"^[A-Z][a-z]*(?:['\-][A-Z]?[a-z]+)*\.?$")
NAME_SKIP_WORDS = {
    "resume", "curriculum vitae", "cv", "profile", "contact", "objective",
    "summary", "career objective", "personal details", "address",
}

SKILL_KEYWORDS = {
    "python", "java", "javascript", "typescript", "c++", "c#", "go", "rust", "ruby",
    "php", "scala", "kotlin", "swift", "sql", "nosql", "mongodb", "postgresql",
    "mysql", "redis", "elasticsearch", "docker", "kubernetes", "aws", "azure",
    "gcp", "terraform", "ansible", "jenkins", "git", "ci/cd", "react", "angular",
    "vue", "node.js", "django", "flask", "fastapi", "spring", "express", "redux",
    "html", "css", "tailwind", "bootstrap", "machine learning", "deep learning",
    "nlp", "tensorflow", "pytorch", "scikit-learn", "pandas", "numpy", "spark",
    "hadoop", "kafka", "rabbitmq", "graphql", "rest", "microservices", "linux",
    "bash", "shell", "agile", "scrum", "jira", "tableau", "power bi", "excel",
    "communication", "leadership", "problem solving", "data analysis", "power bi",
    "etl", "airflow", "snowflake", "databricks", "data engineering", "devops",
    "cybersecurity", "penetration testing", "network security",
}

def extract_skills(text: str) -> List[str]:
    found = set()
    lower = text.lower()
    for skill in SKILL_KEYWORDS:
        if re.search(r"\b" + re.escape(skill) + r"\b", lower):
            found.add(skill)
    if _nlp is not None:
        doc = _nlp(text)
        for ent in doc.ents:
            if ent.label_ == "ORG" and ent.text.lower() in SKILL_KEYWORDS:
                found.add(ent.text.lower())
    return sorted(found)

def extract_email(text: str) -> Optional[str]:
    m = EMAIL_RE.search(text)
    return m.group(0) if m else None

def extract_phone(text: str) -> Optional[str]:
    m = PHONE_RE.search(text)
    return m.group(0) if m else None

def extract_name(text: str) -> Optional[str]:
    """Heuristically pull the candidate's name from the top of the resume.

    Handles the common header styles: Title Case ("John Smith"),
    ALL CAPS ("JOHN SMITH"), and a name line followed by extra content on
    the same line ("John Smith | Software Engineer", "John Smith, Resume").
    """
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    for ln in lines[:15]:
        if EMAIL_RE.fullmatch(ln.strip()) or PHONE_RE.fullmatch(ln.strip()):
            continue

        # Only look at the part before a separator that usually introduces
        # a title/contact info, e.g. "John Smith | Software Engineer".
        candidate = re.split(r"[|,•\u2022:]", ln)[0].strip()
        if any(w in candidate.lower() for w in NAME_SKIP_WORDS):
            continue
        if EMAIL_RE.search(candidate) or any(ch.isdigit() for ch in candidate):
            continue

        words = candidate.split()
        if not (2 <= len(words) <= 4):
            continue
        if not all(NAME_WORD_RE.match(w) for w in words):
            continue

        all_caps = all(w.isupper() for w in words)
        title_case = all(NAME_TITLE_WORD_RE.match(w) for w in words if len(w) > 1)
        if all_caps or title_case:
            return " ".join(w.capitalize() if w.isupper() else w for w in words)
    return None

def extract_experience_years(text: str) -> Optional[int]:
    patterns = [
        r"(\d+)\+?\s*years?\s*(?:of)?\s*experience",
        r"experience\s*(?:of)?\s*(\d+)\+?\s*years?",
        r"(\d+)\+?\s*yrs?",
    ]
    for pat in patterns:
        m = re.search(pat, text.lower())
        if m:
            return int(m.group(1))
    return None

def extract_education(text: str) -> List[str]:
    degrees = ["bachelor", "master", "b.tech", "b.e", "m.tech", "m.e", "phd",
               "mba", "b.sc", "m.sc", "bca", "mca", "diploma"]
    found = []
    lower = text.lower()
    for d in degrees:
        if d in lower:
            found.append(d)
    return found

def extract_certifications(text: str) -> List[str]:
    certs = ["aws certified", "azure certified", "gcp certified", "cka",
             "ckad", "pmp", "scrum master", "cissp", "ceh", "comptia",
             "oracle certified", "tensorflow certified", " Tableau certified", "six sigma"]
    found = []
    lower = text.lower()
    for c in certs:
        if c.strip() in lower:
            found.append(c.strip())
    return found

def extract_projects(text: str) -> List[str]:
    projects = []
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    in_proj = False
    for ln in lines:
        low = ln.lower()
        if "project" in low:
            in_proj = True
            continue
        if in_proj and len(ln) > 5:
            projects.append(ln[:100])
        if len(projects) >= 10:
            break
    return projects

def parse_resume(path: str, candidate_location: Optional[str] = None) -> Dict:
    text = extract_text(path)
    name = extract_name(text)
    email = extract_email(text)
    phone = extract_phone(text)
    skills = extract_skills(text)
    exp_years = extract_experience_years(text)
    education = extract_education(text)
    certs = extract_certifications(text)
    projects = extract_projects(text)
    entities = {
        "name": name,
        "email": email,
        "phone": phone,
        "skills": skills,
        "experience_years": exp_years,
        "education": education,
        "certifications": certs,
        "projects": projects,
        "location": candidate_location,
    }
    return {"raw_text": text, "entities": entities}